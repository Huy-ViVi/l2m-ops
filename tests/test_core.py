"""Test logic lõi bằng file mẫu giả lập.   python tests/test_core.py   (hoặc pytest)"""
from __future__ import annotations

import io
import os
import sys
import tempfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ.setdefault("L2M_DATA_DIR", tempfile.mkdtemp(prefix="l2m-test-"))  # không đụng data/ thật

from openpyxl import load_workbook  # noqa: E402

from core import nc_setup as nc  # noqa: E402
from core import payment_errors as pe  # noqa: E402
from core import webshop_audit as wa  # noqa: E402
from core.master import Master, markets_for  # noqa: E402
from tests import make_samples  # noqa: E402

S = make_samples.OUT
S.mkdir(exist_ok=True)
make_samples.master(S / "SAMPLE All Product Information.xlsx")
make_samples.cs_sheet(S / "SAMPLE CS Escalation.xlsx")
make_samples.mytool_export(S / "SAMPLE JVL2M_SEA-orders.xlsx")
MASTER = Master((S / "SAMPLE All Product Information.xlsx").read_bytes())


def test_master():
    assert not MASTER.warnings, MASTER.warnings
    p = MASTER.packages["web.item153"]
    assert p.name == "Seal Growth Package" and p.price_usd == 19.99
    assert [i.name for i in p.items][:2] == ["Pa'agrio Seal Chest", "Eva Seal Chest"]
    assert p.diamonds == 300 and p.purchase_limit == "2 / Account / Weekly"
    assert MASTER.packages["web.item160"].diamonds is None
    assert MASTER.webshop["web.item153"]["prices"] == {"VN": 499000, "TH": 650, "SG": 26.98, "MY": 88.9, "PH": 1090, "ID": 299000}
    twin = MASTER.get("web.vn.key5")
    assert twin.diamonds == 1000 and "web.sa.key5" in twin.source
    assert MASTER.display_name("web.vn.key5") == "Diamond Pack 1000"
    assert MASTER.timeline and MASTER.timeline[0][0] == 121
    assert markets_for("web.item1") == ["SG", "ID", "PH", "MY", "TH", "VN"]
    assert markets_for("web.sa.key5") == ["SG", "ID", "PH", "MY", "TH"] and markets_for("web.vn.mem2") == ["VN"]


def test_helpers():
    assert pe.server_code("[14050] Erica10") == "E10"
    assert pe.server_code("[14031] Leona01") == "L1"
    assert pe.server_code("[14027] Teon07") == "T7"
    valid, bad = pe.extract_orders("7000000000000000006" + "7000000000000000007")
    assert valid == ["7000000000000000006", "7000000000000000007"] and not bad
    valid, bad = pe.extract_orders("700000000000000004")
    assert not valid and bad == ["700000000000000004"]
    assert pe.parse_account(1000000000000000001) == "1000000000000000001"


def _result(overrides=None):
    cases = pe.read_cs((S / "SAMPLE CS Escalation.xlsx").read_bytes(), "Payment", 2)
    rows = pe.read_mytool_export((S / "SAMPLE JVL2M_SEA-orders.xlsx").read_bytes(), "sea")
    return cases, rows, pe.build_result(cases, rows, MASTER, overrides)


def test_mytool_export_columns():
    _, rows, _ = _result()
    r = rows[0]
    assert r["OrderNumber"] == "8000000000000000001" and r["ParentOrderNumber"] == "7000000000000000001"
    # Price/CurrencyCode phải lấy cột trong nhóm PurchasingInfo, không phải WalletInfo
    assert r["Price"] == "499000" and r["CurrencyCode"] == "VND" and r["ProductID"] == "web.item153"
    assert len(rows) == 8


def test_payment_result():
    cases, _, res = _result()
    assert len(cases) == 6
    assert [l.app_trans_id for l in res.lines] == ["APP-1", "APP-2a", "APP-5", "APP-6", "APP-7"]
    reasons = " | ".join(e["Lý do"] for e in res.excluded)
    assert "IAP" in reasons and "thiếu số" in reasons and "không phải DELIVERY_SERVER_ERROR" in reasons
    assert len(res.extra_errors) == 1 and res.extra_errors[0]["OrderNumber"] == "8000000000000000009"
    assert any("Game Account ID khác MyTool" in w for w in res.warnings)
    assert any("Server CS" in w for w in res.warnings)
    assert res.missing_diamonds() == ["web.item160"]
    groups = res.groups()
    assert [g.nickname for g in groups] == ["VESPER", "เปิดบริสุทธิ์", "Moon"]
    assert groups[0].total_diamonds is None  # thiếu Diamonds của web.item160
    audit = dict((msg.split(" —")[0], ok) for ok, msg in res.audit())
    assert audit["Có đủ Diamonds cho mọi gói"] is False


def test_payment_output():
    _, _, res = _result({"web.item160": 50})
    groups = res.groups()
    assert groups[0].total_diamonds == 650 and groups[2].total_diamonds == 1000
    files = res.csv_files()
    assert set(files) == {"VESPER E10.csv", "เปิดบริสุทธิ์ L1.csv", "Moon E10.csv"}
    assert files["VESPER E10.csv"] == b"1000000000000000001,1\r\n"
    wb = load_workbook(io.BytesIO(pe.write_excel(res)))
    ws = wb["Users List"]
    assert [c.value for c in ws[1]] == ["Game Account ID", "Account Number", "Server", "Nickname",
                                        "Transaction ID (AppTransaction ID)", "Package name/ID", "Quantity", "Total Diamond"]
    assert ws["A1"].fill.fgColor.rgb == "FFE87331" and ws["E1"].fill.fgColor.rgb == "FFA6C9EC"
    assert ws["A1"].font.name == "Aptos Narrow" and ws.freeze_panes == "A2"
    assert ws["B2"].value == "1000000000000000001" and ws["B2"].number_format == "@"
    assert ws["H2"].value == "650" and isinstance(ws["H2"].value, str)
    assert "A2:A4" in {str(r) for r in ws.merged_cells.ranges} and "H2:H4" in {str(r) for r in ws.merged_cells.ranges}
    assert ws["C10"].value == "Compensation Mail" and ws["D10"].value == pe.COMPENSATION_TEXT  # cách 3 dòng sau dòng 6
    wp = wb["Packages List"]
    assert [c.value for c in wp[1]] == ["Packages Name", "Product ID", "Selling Price", "Currency", "Diamonds"]
    assert wp["A1"].fill.fgColor.rgb == "FFFFFF00" and wp["C2"].number_format == "#,##0"
    assert pe.excel_name(date(2026, 10, 26)) == "Payment Error Case 1026.xlsx"


def test_nc_plan():
    catalog = nc.ItemCatalog()
    plan = nc.build_plan(MASTER.packages["web.item153"], catalog)
    assert [i["item_id"] for i in plan["items"]] == [195, 138, 139, 2]
    assert all(i["confirmed"] for i in plan["items"])
    assert plan["store_product_key"] == "com.ncv.l2m.web.item153"
    # Giá setup = giá USD làm tròn LÊN (skill package-creation §3.2): 19.99 -> 20 ở Base, Monetary, Reputation.
    assert plan["price_usd"] == plan["reputation_score"] == 20 and plan["excel_price_usd"] == 19.99
    assert plan["limit"] == {"raw": "2 / Account / Weekly", "amount": 2, "user_type": "Account", "duration": 1, "period": "Week", "unit_code": 5}
    assert nc.validate_plan(plan) == []
    bad = nc.build_plan(MASTER.packages["web.item160"], catalog)
    assert bad["items"][0]["item_id"] is None and bad["limit"] is None
    assert any("chưa có Item ID" in i for i in nc.validate_plan(bad))
    assert nc.parse_limit("1 / Account / Daily")["period"] == "Day"
    assert nc.parse_limit("weird").get("error")


def test_item_usage_suggestion():
    """Gợi ý item theo lịch sử gói: đếm gói đã dùng từng Item ID (thứ tự dòng Item Mapping), chỉ so giữa các item cùng tên."""
    from core.master import MasterPackage, PackageItem

    known = {"greater class card 11 times bound": 10, "greater class skin card 11 times bound": 186,
             "fancy class card 11 times bound old": 79, "diamonds": 2}

    class Catalog:
        def lookup(self, name):
            item_id = known.get(nc.item_key(name))
            return {"id": item_id, "confirmed": item_id is not None}

    def pkg(key, row, *names):
        return MasterPackage(key, key, f"Package {key}", row=row, items=[PackageItem(n, 1) for n in names])

    packages = {p.product_key: p for p in [
        pkg("web.item1", 10, "Greater Class Card (11 Times) (Bound)", "Diamonds"),
        pkg("web.item2", 20, "Greater Class Card (11 Times) (Khóa)", "Greater Class Card (11 Times) (Bound)"),  # 1 gói dù 2 dòng
        *[pkg(f"web.item{i}", 20 + i, "Diamonds") for i in range(3, 40)],
        pkg("web.item40", 100, "Greater Class Skin Card (11 Times) (Bound)", "Mystery Box"),
        pkg("web.item41", 101, "Fancy Class Card (11 Times) (Bound) Old"),
        pkg("web.item42", 102, "Greater Class Card (11 Times) (Bound)"),  # gói đang setup -> không tính
    ]}
    usage, share = nc.item_usage(packages, Catalog(), exclude={"web.item42"})
    assert usage[10] == {"count": 2, "recent": 0, "last": "web.item2", "ago": 39}, usage[10]
    assert usage[186] == {"count": 1, "recent": 1, "last": "web.item40", "ago": 1}
    assert usage[79]["ago"] == 0 and usage[2]["count"] == 38
    assert share == 43 / 44  # chỉ 'Mystery Box' chưa biết ID

    tool = [{"item_id": 10, "name": "Greater Class Card (11 Times) (Bound)"},
            {"item_id": 109, "name": "Greater Class Card (11 Times) (Bound) Nw"},
            {"item_id": 127, "name": "Greater Class Card (11 Times) (Event)"},
            {"item_id": 186, "name": "Greater Class Skin Card (11 Times) (Bound)"},
            {"item_id": 79, "name": "Fancy Class Card (11 Times) (Bound) Old"},
            {"item_id": 200, "name": "Fancy Class Card (11 Times) (Bound)"}]
    greater = nc.rank_tool_items("Greater Class Card (11 Times) (Khóa)", tool)
    # 186 có trong gói gần đây hơn nhưng là item khác tên (Skin) -> vẫn gợi ý 10
    assert nc.suggest_by_usage("Greater Class Card (11 Times) (Khóa)", greater, usage)["item_id"] == 10
    fancy = nc.rank_tool_items("Fancy Class Card (11 Times) (Khóa)", tool)
    assert fancy[0]["item_id"] == 200  # trùng tên nhưng chưa gói nào dùng; bản 'Old' có trong gói mới nhất
    assert nc.suggest_by_usage("Fancy Class Card (11 Times) (Khóa)", fancy, usage)["item_id"] == 79
    assert nc.suggest_by_usage("Greater Class Card (11 Times) (Khóa)", greater, {}) is None

    same = nc.same_item_name
    assert same("Greater Class Card (11 Times) (Khóa)", "Greater Class Card (11 Times) (Bound) Nw")
    assert not same("Greater Class Card (11 Times) (Khóa)", "Greater Class Skin Card (11 Times) (Bound)")
    assert not same("Greater Class Card (11 Times) (Khóa)", "Greater Class Card (11 Times) (Event)")  # khác Bound
    assert not same("Greater Class Card (1 Time) (Khóa)", "Greater Class Card (11 Times) (Bound)")


def test_webshop_audit():
    expected, warns = wa.expected_rows(MASTER, ["web.item153", "web.sa.key5"])
    assert len(expected) == 11 and not warns
    expected = wa.apply_timeline(expected, [{"package_id": "web.item153", "shop": "ALL", "from": "00:00 27/10/2026", "to": "23:59 02/11/2026"}])
    actual = [{"Package ID": e["Package ID"], "Shop": e["Shop"], "Tên gói": e["Tên gói"], "Giá": e["Giá"],
               "Từ": e["Từ"] and "2026-10-27 00:00", "Đến": e["Từ"] and "2026-11-02 23:59:00", "Trạng thái": ""} for e in expected]
    actual[0]["Giá"] = 25  # SG sai giá
    actual[1]["Đến"] = "2026-11-03 23:59"  # ID sai thời gian kết thúc
    actual = [a for a in actual if not (a["Package ID"] == "web.sa.key5" and a["Shop"] == "TH")]  # TH chưa setup
    results = wa.compare(expected, actual)
    errors = [(r["Package ID"], r["Shop"], r["Trường"]) for r in results if r["Kết quả"].startswith("❌")]
    assert errors == [("web.item153", "SG", "Giá bán"), ("web.item153", "ID", "Thời gian đến"), ("web.sa.key5", "TH", "Gói")]
    assert "Tổng kết: 2 gói, 3 lỗi" in wa.report_markdown(results)
    parsed = wa.parse_mytool_row({"Package Name": "Seal", "Selling Price": "26.98", "Package Sale Time": "2026-10-27 00:00 ~ 2026-11-02 23:59", "Status": "on"})
    assert parsed["Giá"] == 26.98 and parsed["Từ"] == "00:00 27/10/2026" and parsed["Đến"] == "23:59 02/11/2026"


def test_master_real_quirks():
    from core.master import key_variants, markets_for, short_key

    p = MASTER.packages["web.item5"]  # key gốc 'com.ncv.l2m.web.item 5'
    assert p.name == "Weekly Top-up Class box"
    assert [(i.name, i.quantity) for i in p.items] == [("Diamonds", 20), ("Greater Class Card (1 Time) (Bound)", 0)]
    assert p.diamonds == 20
    assert MASTER.packages["web.dia.tier1"].diamonds == 40
    last = MASTER.packages["web.item196"]
    assert len(last.items) == 2 and last.diamonds == 40  # bảng phía dưới không bị gộp vào
    assert short_key("com.ncv.l2m.web.item 5") == "web.item5"
    assert key_variants("vn.key1") == ["vn.key1", "web.vn.key1"]
    assert markets_for("com.ncv.l2m.vn.key1") == ["VN"] and markets_for("sa.key2") == ["SG", "ID", "PH", "MY", "TH"]
    assert MASTER.get("item153").product_key == "web.item153"  # MyTool có thể bỏ tiền tố web.
    lim = nc.parse_limit
    assert lim("1 Time / Account / Daily")["period"] == "Day" and lim("2 Times / Account / Weekly")["amount"] == 2
    assert lim("2 / Account / Week")["unit_code"] == 5
    assert lim("No") is None
    one = lim("1 / Account")
    assert one["amount"] == 1 and one["period"] == "Permanent" and one["unit_code"] == 9 and "error" not in one
    assert lim("1 Time / Account")["unit_code"] == 9 and lim("3 / Account / Monthly")["unit_code"] == 6
    long = nc.build_plan(MASTER.packages["web.item153"], nc.ItemCatalog())
    long["name"] = "Einhasad's Shirt Superb Blesssing Packages"  # 42 ký tự > 40 -> phải rút gọn
    assert any("> 40" in i for i in nc.validate_plan(long))
    plan = nc.build_plan(last, nc.ItemCatalog())
    assert not any("Purchase Limit" in i for i in nc.validate_plan(plan))
    assert nc.ItemCatalog().lookup("Relic of Blessings (Khóa)")["id"] == 25


def test_nc_page_usage_hint():
    """Trang Tạo gói NC: ô chọn item ghi số gói đã dùng từng ứng viên, ⭐ item cùng tên có trong gói trước đó và chọn sẵn nó;
    danh sách ứng viên đổi thứ tự thì lựa chọn đã chọn vẫn trỏ đúng item."""
    import json
    import shutil

    from streamlit.testing.v1 import AppTest

    from core import config, store

    os.environ.update(L2M_APP_MODE="local", L2M_NO_AUTOMATION="1")
    d = config.master_dir("l2m")
    d.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(S / "SAMPLE All Product Information.xlsx", d / "current.xlsx")
    (d / "current.json").write_text(json.dumps({"name": "SAMPLE", "uploaded_at": "now"}), encoding="utf-8")
    # Giả lập: gói web.item5 đã dùng item 5003 (dòng 'Greater Class Card (1 Time) (Bound)' được ghi nhận là 5003).
    store.upsert_item(5003, "Unknown Shiny Stone Old", ["Greater Class Card (1 Time) (Bound)"], by="test")
    store.save_tool_items("l2m", [{"item_id": 5001, "name": "Unknown Shiny Stone New"},
                                  {"item_id": 5003, "name": "Unknown Shiny Stone Old"}])
    at = AppTest.from_file(str(ROOT / "views" / "nc_setup.py"), default_timeout=120)
    at.run()
    at.text_area(key="nc_paste_l2m").set_value("web.item160").run()
    assert not at.exception, at.exception
    box = next(s for s in at.selectbox if s.label == "Item trên tool")
    assert box.value == 5003, box.options  # không có item trùng tên, không có AI -> chọn sẵn item ⭐
    star = next(o for o in box.options if "5003" in o)
    assert star.startswith("⭐ 5003") and "1 gói, gần nhất web.item5" in star, star
    assert "chưa gói nào dùng" in next(o for o in box.options if o.startswith("5001"))
    note = " ".join(c.value for c in at.caption)
    assert "⭐ Nên dùng **5003**" in note and "Weekly Top-up Class box (2 gói trước)" in note, note
    assert "chưa gói nào dùng: 5001" in note, note

    box.set_value(5001).run()
    store.save_tool_items("l2m", [{"item_id": 5000, "name": "Unknown Shiny Stone"}])  # ứng viên mới đứng đầu danh sách
    at.run()
    box = next(s for s in at.selectbox if s.label == "Item trên tool")
    assert box.options[1].startswith("5000") and box.value == 5001, (box.options, box.value)


def test_netcheck():
    import socket

    from core import netcheck

    err = ("Page.goto: net::ERR_NAME_NOT_RESOLVED at https://rc-ncv-billingadmin.ncsoft.net/MobileGoodsAdmin/items\n"
           "Call log:\n  - navigating to ...")
    msg = netcheck.explain(err)
    assert "VPN" in msg and "rc-ncv-billingadmin.ncsoft.net" in msg
    assert netcheck.explain("TimeoutError: something else") is None
    bad = netcheck.check("https://khong-ton-tai.invalid/x")  # .invalid luôn NXDOMAIN (RFC 6761)
    assert bad["step"] == "dns" and not bad["ok"]
    assert netcheck.check("")["step"] == "config"
    with socket.socket() as srv:
        srv.bind(("127.0.0.1", 0))
        srv.listen(1)
        assert netcheck.check(f"http://127.0.0.1:{srv.getsockname()[1]}/")["ok"]


if __name__ == "__main__":
    failed = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"PASS {name}")
            except Exception as e:  # noqa: BLE001
                failed += 1
                print(f"FAIL {name}: {type(e).__name__}: {e}")
    sys.exit(1 if failed else 0)
