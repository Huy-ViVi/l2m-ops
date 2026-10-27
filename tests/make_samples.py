"""Tạo file mẫu GIẢ LẬP (không phải dữ liệu thật) theo đúng cấu trúc mô tả trong skill, để test app.

    python tests/make_samples.py   -> tests/samples/*.xlsx
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).with_name("samples")


def master(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Item Mapping"
    ws.append(["Package Name", "Product Key", "Price ($)", "Purchase Limit", "Included Items", "Quantity", "Rebate", "", "Diamonds"])
    blocks = [
        ("Seal Growth Package", "web.item153", 19.99, "2 / Account / Weekly",
         [("Pa'agrio Seal Chest", 6), ("Eva Seal Chest", 6), ("Fancy Soul Crystal (11 Times)", 6), ("Diamond", 300)]),
        ("Earring Growth Package", "web.item154", 19.99, "2 / Account / Weekly",
         [("Fire Dragon's Earring Chest (Bound)", 6), ("Earth Dragon's Earring Chest (Bound)", 6), ("Diamonds", 300)]),
        ("Diamond Pack 1000", "web.sa.key5", 9.99, "1 / Account / Daily", [("Diamonds", 1000)]),
        ("Mystery Box", "web.item160", 4.99, "", [("Unknown Shiny Stone", 3)]),
        # Các kiểu ghi gặp trong master thật: key có dấu cách, tên 2 dòng, số lượng nằm trong tên, limit lạ.
        ("Hot deal\nWeekly Top-up Class box", "com.ncv.l2m.web.item 5", 2.99, "1 Time / Account / Weekly",
         [("Diamonds x 20", None), ("Greater Class Card (1 Time) (Bound)", None)]),
        ("Diamond Packs", "com.ncv.l2m.web.dia.tier1", None, "No", [("40 Diamonds", None)]),
        ("Magic Cannon", "com.ncv.l2m.web.item196", 0.8, "1 / Account", [("Diamonds", 40), ("Parchment of Mana (Bound)", 210)]),
    ]
    for name, key, price, limit, items in blocks:
        first = True
        for item, qty in items:
            ws.append([name if first else None, key if first else None, price if first else None,
                       limit if first else None, item, qty, None, None, None])
            first = False
    # Bên dưới bảng chính, cách vài dòng trống, là một bảng khác — không được gộp vào gói cuối.
    ws.append([])
    ws.append([])
    ws.append(["Gacha rates", None, None, None, "Diamonds", 3000])
    ws.append([None, None, None, None, "Legendary Agathion", 1])

    ws = wb.create_sheet("Webshop")
    ws.append(["Webshop price list"])
    ws.append(["Package Name", "Product Key", "Price USD", "VN", "TH", "Sing", "MY", "PH", "Indo"])
    ws.append(["Hot Deal"])
    ws.append(["Seal Growth Package", "web.item153", 19.99, 499000, 650, 26.98, 88.9, 1090, 299000])
    ws.append(["Earring Growth Package", "web.item154", 19.99, 499000, 650, 26.98, 88.9, 1090, 299000])
    ws.append(["Key"])
    ws.append(["Diamond Pack 1000 SEA", "web.sa.key5", 9.99, None, 329, 13.48, 44.9, 549, 149000])
    ws.append(["Diamond Pack 1000 VN", "web.vn.key5", 9.99, 249000, None, None, None, None, None])

    ws = wb.create_sheet("Webshop Packages Name & Des")
    ws.append(["Product Key", "Name EN", "Name VN", "Description EN"])
    ws.append(["web.item153", "Seal Growth Package", "Gói Phát Triển Ấn", "6x Pa'agrio's Seal Chest ..."])
    ws.append(["web.vn.key5", "Diamond Pack 1000", "Gói 1000 Kim Cương", "1000 Diamonds"])

    ws = wb.create_sheet("Timeline Selling")
    ws.append(["Product", "Product Key", "Start", "End"])
    for _ in range(119):
        ws.append([None])
    ws.append(["Seal Growth Package", "web.item153", "27/10/2026", "02/11/2026"])
    wb.save(path)


def cs_sheet(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Payment"
    ws.append(["Ticket", "Date", "Country", "Game Account ID", "Server", "Account Number", "Email", "Nickname",
               "Note", "Payment provider", "Transaction ID", "Package name", "Purchase date"])
    rows = [
        ["T-001", None, "VN", "GA-AAA", "[14050] Erica10", "1000000000000000001", None, "VESPER", None,
         "Web Shop", "7000000000000000001", "Seal Growth Package", "2026-10-20"],
        # 2 mã trong 1 ô (xuống dòng) + order con lẫn SUCCESS/ERROR
        ["T-002", None, "TH", "GA-BBB", "[14031] Leona01", "1000000000000000002", None, "เปิดบริสุทธิ์", None,
         "Web Shop", "7000000000000000002\n7000000000000000003", "Earring", "2026-10-21"],
        ["T-003", None, "ID", "GA-CCC", "[14027] Teon07", "1000000000000000003", None, "Hiep007", None,
         "GooglePlay", "GPA.1234-5678-9012-34567", "Diamonds", "2026-10-21"],
        # mã thiếu 1 số (18 chữ số)
        ["T-004", None, "PH", "GA-DDD", "[14027] Teon07", "1000000000000000004", None, "Kuro", None,
         "Web Shop", "700000000000000004", "Seal", "2026-10-22"],
        # Game Account ID lệch MyTool + server lệch
        ["T-005", None, "MY", "GA-EEE", "[14050] Erica10", "1000000000000000005", None, "Moon", None,
         "Web Shop", "7000000000000000005", "Diamond Pack", "2026-10-22"],
        # 2 mã bị nối liền (38 chữ số), cùng user VESPER -> gộp nhóm
        ["T-006", None, "VN", "GA-AAA", "[14050] Erica10", "1000000000000000001", None, "VESPER", None,
         "Web Shop", "70000000000000000067000000000000000007", "Mystery", "2026-10-23"],
    ]
    for r in rows:
        ws.append(r)
    wb.save(path)


def mytool_export(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["JVL2M_SEA orders export"])
    top = ["AppTransID", "Region", "ServerID", "UserID", "DeliveryStatus", "OrderNumber", "ParentOrderNumber",
           "CreatedTime", "PurchasingInfo", None, None, None, None, "WalletInfo", None]
    sub = [None] * 8 + ["ProductID", "ProductName", "Price", "Quantity", "CurrencyCode", "Price", "CurrencyCode"]
    ws.append(top)
    ws.append(sub)
    for col in "ABCDEFGH":
        ws.merge_cells(f"{col}2:{col}3")
    ws.merge_cells("I2:M2")
    ws.merge_cells("N2:O2")
    err, ok = "[-1]-[DE]DELIVERY_SERVER_ERROR", "[1]-[DE]DELIVERY_SUCCESS"
    data = [
        ["APP-1", "VN-VND", "14050", "GA-AAA", err, "8000000000000000001", "7000000000000000001", "2026-10-20 10:00", "web.item153", "Seal Growth Pkg", 499000, 1, "VND", 0, "DIA"],
        ["APP-2a", "TH-THB", "14031", "GA-BBB", err, "8000000000000000002", "7000000000000000002", "2026-10-21 11:00", "web.item154", "Earring Growth Package", 650, 1, "THB", 0, "DIA"],
        ["APP-2b", "TH-THB", "14031", "GA-BBB", ok, "8000000000000000022", "7000000000000000002", "2026-10-21 11:00", "web.item154", "Earring Growth Package", 650, 1, "THB", 0, "DIA"],
        ["APP-3", "PH-PHP", "14031", "GA-BBB", ok, "8000000000000000003", "7000000000000000003", "2026-10-21 12:00", "web.item153", "Seal Growth Package", 1090, 1, "PHP", 0, "DIA"],
        ["APP-5", "MY-MYR", "14051", "GA-XXX", err, "8000000000000000005", "7000000000000000005", "2026-10-22 09:00", "web.vn.key5", "Diamond Pack 1000", 249000, 1, "VND", 0, "DIA"],
        ["APP-6", "VN-VND", "14050", "GA-AAA", err, "8000000000000000006", "7000000000000000006", "2026-10-23 09:00", "web.item160", "Mystery Box", 125000, 1, "VND", 0, "DIA"],
        ["APP-7", "VN-VND", "14050", "GA-AAA", err, "8000000000000000007", "7000000000000000007", "2026-10-23 09:05", "web.item153", "Seal Growth Package", 499000, 1, "VND", 0, "DIA"],
        # Order lỗi không có trong CS
        ["APP-9", "SG-SGD", "14010", "GA-ZZZ", err, "8000000000000000009", "7000000000000000009", "2026-10-23 10:00", "web.item154", "Earring Growth Package", 26.98, 1, "SGD", 0, "DIA"],
    ]
    for r in data:
        ws.append(r)
    wb.save(path)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    master(OUT / "SAMPLE All Product Information.xlsx")
    cs_sheet(OUT / "SAMPLE CS Escalation.xlsx")
    mytool_export(OUT / "SAMPLE JVL2M_SEA-orders.xlsx")
    print("Đã tạo file mẫu trong", OUT)
