"""Task 10 — Kiểm tra & tổng hợp case nạp lỗi.

Đầu vào: sheet CS escalation (tab Payment, tải về .xlsx), file Export Excel của MyTool Search Order,
master "All Product Information". Đầu ra: `Payment Error Case <MMDD>.xlsx` + CSV per-user.
Rule: chỉ order có status DELIVERY_SERVER_ERROR mới là lỗi thật; IAP (AppStore/GooglePlay) bị loại.
"""
from __future__ import annotations

import io
import re
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .excel import cell_text, col_index, find_sheet, get, norm, open_workbook, sheet_rows, to_number
from .i18n import ai_language, t
from .master import Master, Msg, short_key, tr  # noqa: F401 — tr dùng ở views (pe.tr)

ORDER_LEN = 19

# Lý do loại, cảnh báo, checklist là Msg: chuỗi tiếng Việt (ghi vào file Payment Error Case, lưu lịch sử —
# file giao CS giống hệt nhau dù xuất ở ngôn ngữ nào), giao diện hiển thị qua tr().

# Cột mặc định của sheet CS "[NCV] - [L2M] Support Escalation (NC)", tab Payment.
CS_COLUMNS = {
    "ticket": "A",
    "country": "C",
    "user_id": "D",
    "server": "E",
    "account": "F",
    "nickname": "H",
    "provider": "J",
    "transactions": "K",
    "note": "L",
    "purchase_date": "M",
}
CS_COLUMN_LABELS = {
    "ticket": "Ticket",
    "country": "Country",
    "user_id": "Game Account ID (UserID)",
    "server": "Server",
    "account": "Account Number",
    "nickname": "Nickname",
    "provider": "Payment provider",
    "transactions": "Transaction ID (Order Number)",
    "note": "Package name / ghi chú",  # nhãn ô nhập trên giao diện -> hiển thị qua t()
    "purchase_date": "Purchase date",
}

MYTOOL_FIELDS = [
    "AppTransID", "Region", "ServerID", "UserID", "DeliveryStatus", "OrderNumber", "ParentOrderNumber",
    "CreatedTime", "ProductID", "ProductName", "Price", "Quantity", "CurrencyCode",
]
IAP_PATTERN = re.compile(r"app\s*store|apple|\bios\b|google|play\s*store|android|\biap\b", re.I)

FONT_NAME = "Aptos Narrow"
FILL_USER = PatternFill("solid", fgColor="FFE87331")
FILL_ORDER = PatternFill("solid", fgColor="FFA6C9EC")
FILL_PACKAGE = PatternFill("solid", fgColor="FFFFFF00")
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
COMPENSATION_TEXT = "This is the compensation Diamonds you received due to an issue with the purchase package."


def _digits(value) -> str:
    return re.sub(r"\D", "", cell_text(value))


def _key(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def is_server_error(status: str) -> bool:
    return "SERVER_ERROR" in (status or "").upper()


def parse_account(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.0f}"
    return re.sub(r"[\s']", "", cell_text(value))


def extract_orders(text: str) -> tuple[list[str], list[str]]:
    """Tách Order Number (luôn 19 chữ số). Trả về (mã hợp lệ, mã sai độ dài)."""
    valid, bad = [], []
    for run in re.findall(r"\d+", text or ""):
        if len(run) == ORDER_LEN:
            valid.append(run)
        elif len(run) > ORDER_LEN and len(run) % ORDER_LEN == 0:
            # Ô nhiều dòng bị nối liền -> tách theo bội số 19.
            valid.extend(run[i:i + ORDER_LEN] for i in range(0, len(run), ORDER_LEN))
        elif len(run) >= 10:
            bad.append(run)
    return list(dict.fromkeys(valid)), bad


def server_code(server: str) -> str:
    """'[14050] Erica10' -> 'E10', '[14031] Leona01' -> 'L1'."""
    name = (server or "").split("]")[-1].strip()
    m = re.match(r"([^\W\d_]+)\s*0*(\d+)", name)
    if not m:
        return name
    return m.group(1)[0].upper() + str(int(m.group(2)))


def safe_filename(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip().rstrip(".") or "unknown"


# ---------------------------------------------------------------------------
# Đọc dữ liệu
# ---------------------------------------------------------------------------
@dataclass
class CSCase:
    row: int
    ticket: str = ""
    country: str = ""
    user_id: str = ""
    server: str = ""
    account: str = ""
    nickname: str = ""
    provider: str = ""
    transactions_raw: str = ""
    note: str = ""
    purchase_date: str = ""
    order_numbers: list[str] = field(default_factory=list)
    bad_tokens: list[str] = field(default_factory=list)

    @property
    def is_iap(self) -> bool:
        p = self.provider or ""
        return bool(IAP_PATTERN.search(p)) and "web" not in p.lower()


def list_sheets(src) -> list[str]:
    wb = open_workbook(src)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_cs(src, sheet: str = "Payment", start_row: int = 2, end_row: int | None = None,
            columns: dict[str, str] | None = None) -> list[CSCase]:
    columns = columns or CS_COLUMNS
    wb = open_workbook(src)
    try:
        ws = wb[sheet] if sheet in wb.sheetnames else (find_sheet(wb, sheet) or wb.worksheets[0])
        idx = {k: col_index(v) for k, v in columns.items() if v}
        cases: list[CSCase] = []
        for r, raw in enumerate(ws.iter_rows(min_row=start_row, max_row=end_row, values_only=True), start=start_row):
            row = list(raw)

            def g(k):
                return get(row, idx.get(k))

            trans = g("transactions")
            case = CSCase(
                row=r, ticket=cell_text(g("ticket")), country=cell_text(g("country")),
                user_id=cell_text(g("user_id")), server=cell_text(g("server")), account=parse_account(g("account")),
                nickname=cell_text(g("nickname")), provider=cell_text(g("provider")),
                transactions_raw=cell_text(trans), note=cell_text(g("note")), purchase_date=cell_text(g("purchase_date")),
            )
            if not (case.user_id or case.account or case.transactions_raw):
                continue
            if isinstance(trans, float):
                case.bad_tokens.append(Msg.t("{raw} (ô dạng số — Excel đã làm tròn mã)", raw=case.transactions_raw))
            else:
                case.order_numbers, case.bad_tokens = extract_orders(case.transactions_raw)
            cases.append(case)
        return cases
    finally:
        wb.close()


def read_mytool_export(src, label: str = "") -> list[dict]:
    """Đọc file `<AppAlias>-orders.xlsx`: header nhóm ở dòng 2, cột con ở dòng 3, data từ dòng 4.
    Dò cột theo tên chứ không theo vị trí."""
    wb = open_workbook(src)
    try:
        rows = sheet_rows(wb.worksheets[0])
    finally:
        wb.close()
    field_by_key = {_key(f): f for f in MYTOOL_FIELDS}
    header_rows = [i for i, row in enumerate(rows[:8]) if any(_key(c) in field_by_key for c in row)]
    if not header_rows:
        raise ValueError(t("Không nhận diện được file export MyTool (không thấy cột OrderNumber / DeliveryStatus)."))
    top, bottom = header_rows[0], header_rows[-1]
    ncols = max(len(r) for r in rows[top:bottom + 1])
    candidates: dict[str, list[tuple[int, str]]] = {}
    group = ""
    for c in range(ncols):
        top_val = cell_text(get(rows[top], c))
        if top_val:
            group = top_val
        name = ""
        for hr in range(bottom, top - 1, -1):
            name = cell_text(get(rows[hr], c))
            if name:
                break
        f = field_by_key.get(_key(name))
        if f:
            candidates.setdefault(f, []).append((c, group if name != top_val else ""))
    col_of = {}
    for f, opts in candidates.items():
        preferred = [c for c, g in opts if "purchas" in g.lower()]
        col_of[f] = preferred[0] if preferred else opts[0][0]
    if "OrderNumber" not in col_of or "DeliveryStatus" not in col_of:
        raise ValueError(t("File export MyTool thiếu cột OrderNumber hoặc DeliveryStatus."))
    records = []
    for row in rows[bottom + 1:]:
        rec = {f: cell_text(get(row, c)) for f, c in col_of.items()}
        if not rec.get("OrderNumber"):
            continue
        rec["_source"] = label
        records.append(rec)
    return records


# ---------------------------------------------------------------------------
# Đối chiếu
# ---------------------------------------------------------------------------
@dataclass
class ErrorLine:
    case: CSCase
    user_id: str
    app_trans_id: str
    order_number: str
    parent_order: str
    product_id: str
    package_name: str
    mytool_product_name: str
    quantity: int
    price: float | None
    currency: str
    diamonds: int | None
    status: str
    created: str
    region: str
    mytool_server: str

    @property
    def total_diamonds(self) -> int | None:
        return None if self.diamonds is None else self.diamonds * self.quantity


@dataclass
class UserGroup:
    user_id: str
    account: str
    server: str
    nickname: str
    lines: list[ErrorLine]

    @property
    def total_diamonds(self) -> int | None:
        values = [l.total_diamonds for l in self.lines]
        return None if any(v is None for v in values) else sum(values)

    @property
    def csv_name(self) -> str:
        return safe_filename(f"{self.nickname} {server_code(self.server)}") + ".csv"

    @property
    def csv_bytes(self) -> bytes:
        return f"{self.account},1\r\n".encode("utf-8")


@dataclass
class Result:
    cases: list[CSCase]
    lines: list[ErrorLine] = field(default_factory=list)
    excluded: list[dict] = field(default_factory=list)  # "Lý do" (và "Order" nếu là mã sai) là Msg
    warnings: list[Msg] = field(default_factory=list)
    extra_errors: list[dict] = field(default_factory=list)
    master_notes: list[Msg] = field(default_factory=list)

    def groups(self) -> list[UserGroup]:
        out: dict[tuple, UserGroup] = {}
        for line in self.lines:
            c = line.case
            k = (line.user_id, c.account, c.server, c.nickname)
            if k not in out:
                out[k] = UserGroup(line.user_id, c.account, c.server, c.nickname, [])
            out[k].lines.append(line)
        return list(out.values())

    def packages(self) -> list[dict]:
        seen: dict[tuple, dict] = {}
        for line in self.lines:
            k = (line.product_id, line.currency)
            if k not in seen:
                seen[k] = {
                    "Packages Name": line.package_name, "Product ID": line.product_id,
                    "Selling Price": line.price, "Currency": line.currency, "Diamonds": line.diamonds,
                }
        return list(seen.values())

    def missing_diamonds(self) -> list[str]:
        return sorted({l.product_id for l in self.lines if l.diamonds is None})

    def csv_files(self) -> dict[str, bytes]:
        files: dict[str, bytes] = {}
        for g in self.groups():
            name = g.csv_name
            n = 2
            while name in files:
                name = g.csv_name[:-4] + f" ({n}).csv"
                n += 1
            files[name] = g.csv_bytes
        return files

    def audit(self) -> list[tuple[bool, Msg]]:
        checks: list[tuple[bool, Msg]] = []
        ids = [l.app_trans_id for l in self.lines]
        checks.append((all(ids) and len(ids) == len(set(ids)),
                       Msg.t("AppTransaction ID đầy đủ và không trùng ({unique}/{total})", unique=len(set(ids)), total=len(ids))))
        bad_acc = [g.nickname for g in self.groups() if not re.fullmatch(r"\d{19}", g.account or "")]
        checks.append((not bad_acc, Msg.t("Account Number đủ 19 số — sai: {names}", names=", ".join(bad_acc)) if bad_acc
                       else Msg.t("Account Number đủ 19 số")))
        mismatch = [w for w in self.warnings if "Game Account ID" in w]
        checks.append((not mismatch, Msg.t("Game Account ID MyTool khớp cột D của CS — {n} lệch", n=len(mismatch)) if mismatch
                       else Msg.t("Game Account ID MyTool khớp cột D của CS")))
        missing = self.missing_diamonds()
        checks.append((not missing, Msg.t("Có đủ Diamonds cho mọi gói — thiếu: {packages}", packages=", ".join(missing)) if missing
                       else Msg.t("Có đủ Diamonds cho mọi gói")))
        csv_ok = all(re.fullmatch(rb"\d{19},1\r\n", b) for b in self.csv_files().values())
        checks.append((csv_ok, Msg.t("CSV đúng dạng `<Account Number>,1` + CRLF")))
        return checks


def build_result(cases: list[CSCase], mytool_rows: list[dict], master: Master | None,
                 diamonds_override: dict[str, int] | None = None) -> Result:
    diamonds_override = {short_key(k): v for k, v in (diamonds_override or {}).items() if v not in (None, "")}
    res = Result(cases=cases)

    by_order: dict[str, dict] = {}
    by_parent: dict[str, list[dict]] = defaultdict(list)
    for rec in mytool_rows:
        on = _digits(rec.get("OrderNumber"))
        if not on or on in by_order:
            continue
        by_order[on] = rec
        po = _digits(rec.get("ParentOrderNumber"))
        if po:
            by_parent[po].append(rec)
    short_orders = [on for on in by_order if len(on) != ORDER_LEN]
    if short_orders:
        res.warnings.append(Msg.t("File MyTool có {n} OrderNumber không đủ 19 số (có thể bị Excel làm tròn).", n=len(short_orders)))

    def exclude(case: CSCase, order: str, status: str, reason: Msg):
        res.excluded.append({
            "Dòng CS": case.row, "Ticket": case.ticket, "Nickname": case.nickname,
            "Order": order, "Trạng thái": status, "Lý do": reason,
        })

    referenced: set[str] = set()
    seen: set[str] = set()
    master_cache: dict[str, tuple] = {}

    def package_info(pid: str, mytool_name: str) -> tuple[str, int | None]:
        k = short_key(pid)
        if k not in master_cache:
            pkg = master.get(k) if master else None
            name = (master.display_name(k) if master else "") or (pkg.name if pkg else "") or mytool_name
            dia = pkg.diamonds if pkg else None
            if pkg is None and master is not None:
                res.master_notes.append(Msg.t("{pid}: không có trong master — cần tải bản master mới nhất", pid=pid))
            elif pkg is not None and pkg.note:
                res.master_notes.append(Msg("{pid}: {note}", pid=pid, note=pkg.note))
            if mytool_name and name and norm(mytool_name) != norm(name):
                res.master_notes.append(Msg.t("{pid}: tên master '{name}' khác tên MyTool '{mytool}'",
                                              pid=pid, name=name, mytool=mytool_name))
            master_cache[k] = (name, dia)
        name, dia = master_cache[k]
        if k in diamonds_override:
            dia = int(diamonds_override[k])
        return name, dia

    for case in cases:
        if case.is_iap:
            exclude(case, case.transactions_raw, "",
                    Msg.t("{provider}: IAP không có trên MyTool webpayment — user tự kiểm tra", provider=case.provider))
            continue
        for tok in case.bad_tokens:
            exclude(case, tok, "", Msg.t("Transaction ID sai/thiếu số — tra theo User ID trên MyTool rồi báo CS sửa sheet"))
            res.warnings.append(Msg.t("Dòng CS {row} ({nickname}): mã '{token}' không đủ 19 số",
                                      row=case.row, nickname=case.nickname, token=tok))
        if not case.order_numbers and not case.bad_tokens:
            exclude(case, "", "", Msg.t("Không có Transaction ID"))
        for on in case.order_numbers:
            children = list(by_parent.get(on, []))
            if not children and on in by_order:
                children = [by_order[on]]
            if not children:
                exclude(case, on, "", Msg.t("Không có trong file export MyTool (thử shop còn lại SEA ↔ VN)"))
                continue
            referenced.add(on)
            skipped = Counter()
            for rec in children:
                child_on = _digits(rec.get("OrderNumber"))
                referenced.add(child_on)
                status = rec.get("DeliveryStatus", "")
                if not is_server_error(status):
                    skipped[status or "(trống)"] += 1
                    continue
                uniq = rec.get("AppTransID") or child_on
                if uniq in seen:
                    res.warnings.append(Msg.t("Order {order} xuất hiện ở nhiều dòng CS — chỉ tính 1 lần (dòng {row})",
                                              order=child_on, row=case.row))
                    continue
                seen.add(uniq)
                pid = rec.get("ProductID", "")
                name, dia = package_info(pid, rec.get("ProductName", ""))
                qty = int(to_number(rec.get("Quantity")) or 1)
                mt_user = rec.get("UserID", "")
                if mt_user and case.user_id and norm(mt_user) != norm(case.user_id):
                    res.warnings.append(Msg.t(
                        "Dòng CS {row}: Game Account ID khác MyTool ({cs_user} ≠ {mytool_user}) — order {order}",
                        row=case.row, cs_user=case.user_id, mytool_user=mt_user, order=child_on))
                cs_srv = re.search(r"\[(\d+)\]", case.server or "")
                mt_srv = _digits(rec.get("ServerID"))
                if cs_srv and mt_srv and cs_srv.group(1) != mt_srv:
                    res.warnings.append(Msg.t(
                        "Dòng CS {row}: Server CS {cs_server} ≠ MyTool {mytool_server} — đang giữ theo CS, hỏi lại nếu cần",
                        row=case.row, cs_server=case.server, mytool_server=mt_srv))
                res.lines.append(ErrorLine(
                    case=case, user_id=mt_user or case.user_id, app_trans_id=rec.get("AppTransID", ""),
                    order_number=child_on, parent_order=_digits(rec.get("ParentOrderNumber")), product_id=pid,
                    package_name=name, mytool_product_name=rec.get("ProductName", ""), quantity=qty,
                    price=to_number(rec.get("Price")), currency=rec.get("CurrencyCode", ""), diamonds=dia,
                    status=status, created=rec.get("CreatedTime", ""), region=rec.get("Region", ""), mytool_server=mt_srv,
                ))
            for status, n in skipped.items():
                exclude(case, on, status, Msg.t("{n} order không phải DELIVERY_SERVER_ERROR — loại", n=n))

    for on, rec in by_order.items():
        if is_server_error(rec.get("DeliveryStatus", "")) and on not in referenced \
                and _digits(rec.get("ParentOrderNumber")) not in referenced:
            res.extra_errors.append({
                "OrderNumber": on, "ParentOrderNumber": _digits(rec.get("ParentOrderNumber")),
                "UserID": rec.get("UserID", ""), "ProductID": rec.get("ProductID", ""),
                "CreatedTime": rec.get("CreatedTime", ""), "DeliveryStatus": rec.get("DeliveryStatus", ""),
            })
    if res.extra_errors:
        res.warnings.append(Msg.t(
            "MyTool có thêm {n} order lỗi không nằm trong danh sách CS — không tự thêm, cần hỏi lại.", n=len(res.extra_errors)))
    res.master_notes = list(dict.fromkeys(res.master_notes))
    return res


# ---------------------------------------------------------------------------
# Xuất file
# ---------------------------------------------------------------------------
def _clean(value):
    """Ô từ Google Sheet có thể chứa ký tự điều khiển ẩn -> openpyxl từ chối ghi (IllegalCharacterError)."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def _style(cell, fill=None, bold=False, num_fmt=None, wrap=False):
    cell.font = Font(name=FONT_NAME, size=11, bold=bold)
    cell.border = BORDER
    cell.alignment = Alignment(vertical="center", wrap_text=wrap)
    if fill:
        cell.fill = fill
    if num_fmt:
        cell.number_format = num_fmt


def write_excel(res: Result) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Users List"
    headers = [
        ("Game Account ID", FILL_USER), ("Account Number", FILL_USER), ("Server", FILL_USER), ("Nickname", FILL_USER),
        ("Transaction ID (AppTransaction ID)", FILL_ORDER), ("Package name/ID", FILL_ORDER), ("Quantity", FILL_ORDER),
        ("Total Diamond", FILL_ORDER),
    ]
    for c, (title, fill) in enumerate(headers, start=1):
        _style(ws.cell(row=1, column=c, value=title), fill=fill, bold=True)
    for letter, width in zip("ABCDEFGH", [48.57, 26.14, 22.14, 18.29, 47.86, 28, 9, 16]):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"

    r = 2
    for g in res.groups():
        first = r
        total = g.total_diamonds
        for i, line in enumerate(g.lines):
            values = [
                g.user_id if i == 0 else None, g.account if i == 0 else None,
                g.server if i == 0 else None, g.nickname if i == 0 else None,
                line.app_trans_id, line.product_id, line.quantity,
                (str(total) if total is not None else "") if i == 0 else None,
            ]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=r, column=c, value=_clean(v))
                _style(cell, num_fmt="@" if c == 2 else None)
            r += 1
        if r - first > 1:
            for col in "ABCDH":
                ws.merge_cells(f"{col}{first}:{col}{r - 1}")
    last = r - 1

    mail_row = last + 4
    ws.cell(row=mail_row, column=3, value="Compensation Mail").font = Font(name=FONT_NAME, size=11, bold=True)
    ws.cell(row=mail_row, column=4, value=COMPENSATION_TEXT).font = Font(name=FONT_NAME, size=11)

    notes = [f"Nguồn: sheet CS ({len(res.cases)} dòng) + MyTool Search Order (Export Excel) + master All Product Information."]
    for ex in res.excluded:
        notes.append(f"Loại — dòng CS {ex['Dòng CS']} {ex['Nickname']} {ex['Order']}: {ex['Lý do']}"
                     + (f" [{ex['Trạng thái']}]" if ex["Trạng thái"] else ""))
    notes.extend(f"Lưu ý: {w}" for w in res.warnings)
    notes.extend(f"Master: {n}" for n in res.master_notes)
    row = mail_row + 2
    ws.cell(row=row, column=1, value="Ghi chú").font = Font(name=FONT_NAME, size=11, bold=True)
    for n in notes:
        row += 1
        ws.cell(row=row, column=1, value=_clean(n)).font = Font(name=FONT_NAME, size=11)

    wp = wb.create_sheet("Packages List")
    for c, title in enumerate(["Packages Name", "Product ID", "Selling Price", "Currency", "Diamonds"], start=1):
        _style(wp.cell(row=1, column=c, value=title), fill=FILL_PACKAGE, bold=True)
    for letter, width in zip("ABCDE", [40, 28, 14, 10.29, 16]):
        wp.column_dimensions[letter].width = width
    wp.freeze_panes = "A2"
    for r, p in enumerate(res.packages(), start=2):
        values = [p["Packages Name"], p["Product ID"], p["Selling Price"], p["Currency"], p["Diamonds"]]
        for c, v in enumerate(values, start=1):
            _style(wp.cell(row=r, column=c, value=_clean(v)), num_fmt="#,##0" if c == 3 else None)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_name(day: date) -> str:
    return f"Payment Error Case {day:%m%d}.xlsx"


def build_zip(res: Result, day: date) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(excel_name(day), write_excel(res))
        for name, data in res.csv_files().items():
            z.writestr(name, data)
    return buf.getvalue()


def output_folder(root: str | Path, day: date, game: str = "l2m") -> Path:
    """Lineage2M giữ nguyên Payment-Error-Cases\\<ngày>\\ như trước; game khác có thư mục con riêng."""
    base = Path(root) / "Payment-Error-Cases"
    return (base if game == "l2m" else base / game) / f"{day:%Y-%m-%d}"


def planned_files(res: Result, day: date) -> list[str]:
    return [excel_name(day)] + list(res.csv_files())


def save_to_folder(res: Result, root: str | Path, day: date, game: str = "l2m") -> list[Path]:
    folder = output_folder(root, day, game)
    folder.mkdir(parents=True, exist_ok=True)
    written = []
    path = folder / excel_name(day)
    path.write_bytes(write_excel(res))
    written.append(path)
    for name, data in res.csv_files().items():
        p = folder / name
        p.write_bytes(data)
        written.append(p)
    return written


def summary_system(game_name: str = "Lineage2M SEA") -> str:
    return f"Bạn là trợ lý vận hành game {game_name}, viết báo cáo nội bộ ngắn gọn bằng {ai_language()}."


def summary_prompt(res: Result, game_name: str = "Lineage2M SEA") -> str:
    """Dữ liệu gửi AI giữ tiếng Việt; bản tóm tắt viết theo ngôn ngữ giao diện (ai_language)."""
    groups = res.groups()
    lines = [f"- {g.nickname} ({g.server}): {len(g.lines)} order lỗi, tổng {g.total_diamonds} Diamonds" for g in groups]
    return (
        f"Tóm tắt đợt xử lý case nạp lỗi {game_name} để báo cáo nội bộ ({ai_language()}, ngắn gọn, dạng bullet).\n"
        f"Số dòng CS: {len(res.cases)}; số order lỗi thật: {len(res.lines)}; số user cần đền bù: {len(groups)}.\n"
        "Danh sách user:\n" + "\n".join(lines) + "\n"
        f"Case bị loại ({len(res.excluded)}):\n" + "\n".join(f"- {e['Nickname']} {e['Order']}: {e['Lý do']}" for e in res.excluded) + "\n"
        "Cảnh báo:\n" + "\n".join(f"- {w}" for w in res.warnings + res.master_notes)
        + "\nNêu rõ việc cần làm tiếp (vd: báo CS sửa mã sai, hỏi lại order dư)."
    )
