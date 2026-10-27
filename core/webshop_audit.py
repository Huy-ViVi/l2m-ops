"""Task 4 — Rà soát gói trên webshop VNG MyTool so với Excel master.

Kỳ vọng (expected) lấy từ master: giá theo thị trường (sheet Webshop), tên (Name & Des), thời gian bán
(Timeline Selling — AI hỗ trợ đọc, người dùng duyệt lại). Thực tế (actual) đọc từ MyTool Sold Package.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import config
from .excel import cell_text, norm, to_number
from .i18n import ai_language, t, tr_value
from .master import Master, key_variants, markets_for, short_key

# Khoá và nhãn trong kết quả giữ tiếng Việt (dữ liệu lưu lịch sử, so sánh trong code);
# common.df() / report_* dịch khi hiển thị hoặc xuất báo cáo.
TIME_FIELDS = {"Từ": "Thời gian từ", "Đến": "Thời gian đến"}

DT_FORMATS = [
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d",
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
    "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y/%m/%d",
    "%H:%M:%S %d/%m/%Y", "%H:%M %d/%m/%Y", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
]
DT_REGEX = re.compile(
    r"\d{4}[-/]\d{1,2}[-/]\d{1,2}(?:[ T]\d{1,2}:\d{2}(?::\d{2})?)?"
    r"|\d{1,2}:\d{2}(?::\d{2})?\s+\d{1,2}/\d{1,2}/\d{4}"
    r"|\d{1,2}[/-]\d{1,2}[/-]\d{4}(?:\s+\d{1,2}:\d{2}(?::\d{2})?)?"
)


def parse_dt(value) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    s = re.sub(r"\s+", " ", str(value).strip())
    for fmt in DT_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    found = DT_REGEX.search(s)
    if found and found.group() != s:
        return parse_dt(found.group())
    return None


def find_datetimes(text: str) -> list[datetime]:
    return [d for d in (parse_dt(m.group()) for m in DT_REGEX.finditer(text or "")) if d]


def fmt_dt(d: datetime | None) -> str:
    return d.strftime("%H:%M %d/%m/%Y") if d else ""


def expected_rows(master: Master, package_ids: list[str]) -> tuple[list[dict], list[str]]:
    rows, warnings = [], []
    for pid in package_ids:
        k = short_key(pid)
        if not k:
            continue
        ws = next((master.webshop[v] for v in key_variants(k) if v in master.webshop), None)
        if ws is None:
            warnings.append(t("{key}: không có trong sheet Webshop — giá kỳ vọng để trống", key=k))
        for market in markets_for(k):
            rows.append({
                "Package ID": k, "Shop": market, "Tên gói": master.display_name(k),
                "Giá": (ws or {}).get("prices", {}).get(market), "Từ": "", "Đến": "", "Trạng thái": "",
            })
    return rows, warnings


def parse_mytool_row(row: dict) -> dict:
    """Một dòng kết quả Sold Package (dict header -> text) -> các trường chuẩn."""
    out = {"Tên gói": "", "Giá": None, "Từ": "", "Đến": "", "Trạng thái": ""}
    dts: list[datetime] = []
    for k, v in row.items():
        nk, text = norm(k), cell_text(v)
        if not text:
            continue
        if "name" in nk and not out["Tên gói"]:
            out["Tên gói"] = text
        elif "price" in nk and out["Giá"] is None:
            out["Giá"] = to_number(text)
        elif "status" in nk and not out["Trạng thái"]:
            out["Trạng thái"] = text
        elif nk in ("from", "start", "start time", "sale from") or nk.endswith(" from"):
            out["Từ"] = text
        elif nk in ("to", "end", "end time", "sale to") or nk.endswith(" to"):
            out["Đến"] = text
        elif "time" in nk or "date" in nk:
            dts.extend(find_datetimes(text))
    if dts and not out["Từ"]:
        out["Từ"] = fmt_dt(dts[0])
    if len(dts) > 1 and not out["Đến"]:
        out["Đến"] = fmt_dt(dts[1])
    return out


def compare(expected: list[dict], actual: list[dict]) -> list[dict]:
    """Mỗi phần tử kết quả: Package ID, Shop, Trường, Excel master, MyTool, Kết quả."""
    act = {(short_key(a.get("Package ID")), a.get("Shop")): a for a in actual}
    results = []
    for e in expected:
        k = (short_key(e["Package ID"]), e["Shop"])
        a = act.get(k)
        base = {"Package ID": k[0], "Shop": k[1], "Tên gói": e.get("Tên gói", "")}
        if a is None or a.get("_not_found"):
            results.append({**base, "Trường": "Gói", "Excel master": "có bán", "MyTool": "không tìm thấy", "Kết quả": "❌ SAI"})
            continue
        checks = []
        if e.get("Giá") not in (None, ""):
            ev, av = to_number(e["Giá"]), to_number(a.get("Giá"))
            checks.append(("Giá bán", ev, av, av is not None and ev is not None and abs(ev - av) < 0.5))
        for field, label in TIME_FIELDS.items():
            if e.get(field):
                ed, ad = parse_dt(e[field]), parse_dt(a.get(field))
                ok = ed is not None and ad is not None and ed.replace(second=0) == ad.replace(second=0)
                checks.append((label, fmt_dt(ed) or e[field], fmt_dt(ad) or a.get(field, ""), ok))
        if e.get("Tên gói") and a.get("Tên gói"):
            checks.append(("Tên gói", e["Tên gói"], a["Tên gói"], norm(e["Tên gói"]) == norm(a["Tên gói"])))
        if e.get("Trạng thái"):
            checks.append(("Trạng thái", e["Trạng thái"], a.get("Trạng thái", ""), norm(e["Trạng thái"]) == norm(a.get("Trạng thái"))))
        for field, ev, av, ok in checks:
            results.append({**base, "Trường": field, "Excel master": cell_text(ev), "MyTool": cell_text(av),
                            "Kết quả": "✅ OK" if ok else "❌ SAI"})
    return results


def report_markdown(results: list[dict], day: date | None = None, game_name: str = "Lineage2M") -> str:
    """Báo cáo text (hiện trên trang + tải về) theo ngôn ngữ giao diện."""
    day = day or config.today()
    packages = list(dict.fromkeys(r["Package ID"] for r in results))
    shops = {(r["Package ID"], r["Shop"]) for r in results}
    errors = [r for r in results if r["Kết quả"].startswith("❌")]
    lines = [
        t("📋 Báo cáo Audit Gói Nạp — {game} Webshop", game=game_name),
        t("Ngày kiểm tra: {day}", day=f"{day:%d/%m/%Y}"),
        t("Phạm vi: {packages} gói × {shops} lượt shop", packages=len(packages), shops=len(shops)),
        "",
    ]
    header = ["Shop", "Trường", "Excel master", "MyTool", "Kết quả"]
    for pid in packages:
        rows = [r for r in results if r["Package ID"] == pid]
        name = next((r["Tên gói"] for r in rows if r["Tên gói"]), "")
        lines.append(f"🔹 {pid} — {name}")
        lines.append("| " + " | ".join(tr_value(h) for h in header) + " |")
        lines.append("|" + "|".join("-" * (len(tr_value(h)) + 2) for h in header) + "|")
        for r in rows:
            lines.append("| " + " | ".join(str(tr_value(r[h])) for h in header) + " |")
        lines.append("")
    lines.append(t("Tổng kết: {packages} gói, {errors} lỗi — cần sửa trên MyTool.", packages=len(packages), errors=len(errors))
                 if errors else t("Tổng kết: {packages} gói, {errors} lỗi — tất cả khớp.", packages=len(packages), errors=0))
    return "\n".join(lines)


def report_excel(results: list[dict]) -> bytes:
    """Báo cáo Excel nội bộ của team (không phải file giao bên ngoài) -> tiêu đề/nhãn theo ngôn ngữ giao diện."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit"
    headers = ["Package ID", "Shop", "Tên gói", "Trường", "Excel master", "MyTool", "Kết quả"]
    ws.append([tr_value(h) for h in headers])
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="FFA6C9EC")
    red = PatternFill("solid", fgColor="FFFFC7CE")
    for r in results:
        ws.append([tr_value(r.get(h, "")) for h in headers])
        if r["Kết quả"].startswith("❌"):
            for c in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=c).fill = red
    for letter, width in zip("ABCDEFG", [22, 8, 36, 18, 28, 28, 12]):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


TIMELINE_SCHEMA = {
    "type": "object",
    "properties": {
        "entries": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "package_id": {"type": "string"},
                    "shop": {"type": "string", "enum": ["ALL", "SG", "ID", "PH", "MY", "TH", "VN"]},
                    "from": {"type": "string"},
                    "to": {"type": "string"},
                    "source_rows": {"type": "string"},
                    "note": {"type": "string"},
                },
                "required": ["package_id", "shop", "from", "to", "source_rows", "note"],
                "additionalProperties": False,
            },
        }
    },
    "required": ["entries"],
    "additionalProperties": False,
}


def ai_extract_timeline(ai, master: Master, package_ids: list[str], max_rows: int = 400,
                        game_name: str = "Lineage2M SEA") -> list[dict]:
    keys = [short_key(p) for p in package_ids if short_key(p)]
    names = {k: master.display_name(k) for k in keys}
    header = " | ".join(cell_text(c) for c in master.timeline_header)
    body = "\n".join(
        f"Dòng {idx}: " + " | ".join(cell_text(c) for c in row)
        for idx, row in master.timeline[:max_rows]
    )
    system = (
        f"Bạn đọc sheet 'Timeline Selling' trong file master {game_name} để lấy thời gian bán của gói webshop. "
        "Chỉ dùng dữ liệu được cung cấp, không đoán. Thời gian trả về dạng 'HH:MM DD/MM/YYYY' (giờ GMT+7). "
        "Nếu sheet không ghi giờ: bắt đầu 00:00, kết thúc 23:59. Nếu không tìm thấy gói, trả from/to rỗng và ghi note. "
        "shop = ALL nếu áp dụng mọi shop. "
        f"Viết note bằng {ai_language()}."
    )
    prompt = (
        f"Hôm nay: {config.today():%d/%m/%Y}.\n"
        "Gói cần tìm (Product Key — tên):\n" + "\n".join(f"- {k} — {names[k]}" for k in keys)
        + f"\n\nHeader sheet: {header}\n\nDữ liệu từ dòng {master.timeline[0][0] if master.timeline else '?'}:\n{body}"
    )
    data = ai.json(system, prompt, TIMELINE_SCHEMA, effort="medium")
    return data.get("entries", [])


def apply_timeline(expected: list[dict], entries: list[dict]) -> list[dict]:
    for e in expected:
        for entry in entries:
            if short_key(entry.get("package_id")) == e["Package ID"] and entry.get("shop") in ("ALL", e["Shop"]):
                e["Từ"], e["Đến"] = entry.get("from", ""), entry.get("to", "")
    return expected
