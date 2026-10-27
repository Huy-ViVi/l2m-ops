"""Tiện ích đọc Excel: dò header theo tên cột (không hardcode index) vì vị trí cột trong master hay xê dịch."""
from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


def norm(value) -> str:
    """Chuẩn hoá text để so khớp: thường hoá, bỏ dấu câu, gộp khoảng trắng. Giữ ký tự unicode (VN/TH)."""
    if value is None:
        return ""
    s = str(value).replace(" ", " ").strip().lower()
    s = re.sub(r"['’`]", "", s)
    s = re.sub(r"[^0-9a-zÀ-￿]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def tokens(value) -> set[str]:
    return set(norm(value).split())


def cell_text(value) -> str:
    """Giá trị ô -> chuỗi. Số nguyên không bị thêm '.0', ngày giờ ra ISO."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def to_number(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    s = re.sub(r"[^0-9.,\-]", "", str(value))
    if not s:
        return None
    # "1,234,000" / "1.234.000" / "12.99" / "12,99"
    if s.count(",") > 1 or (s.count(",") == 1 and len(s.split(",")[1]) == 3 and "." not in s):
        s = s.replace(",", "")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3 and "," not in s):
        s = s.replace(".", "")
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def open_workbook(src, read_only: bool = True):
    """src: đường dẫn, bytes, hoặc file upload của Streamlit."""
    if isinstance(src, (bytes, bytearray)):
        src = BytesIO(src)
    elif hasattr(src, "getvalue"):
        src = BytesIO(src.getvalue())
    return load_workbook(src, data_only=True, read_only=read_only)


def sheet_rows(ws, min_row: int = 1, max_row: int | None = None) -> list[list]:
    return [list(r) for r in ws.iter_rows(min_row=min_row, max_row=max_row, values_only=True)]


def find_sheet(wb, *names):
    wanted = [norm(n) for n in names]
    by_norm = {norm(ws_name): ws_name for ws_name in wb.sheetnames}
    for w in wanted:
        if w in by_norm:
            return wb[by_norm[w]]
    for w in wanted:
        for n, real in by_norm.items():
            if w and w in n:
                return wb[real]
    return None


def find_header_row(rows: list[list], must_have: list[str], scan: int = 15) -> int | None:
    keys = [norm(k) for k in must_have]
    for i, row in enumerate(rows[:scan]):
        cells = [norm(c) for c in row]
        if all(any(k in c for c in cells) for k in keys):
            return i
    return None


def find_col(header_row: list, *candidates: str) -> int | None:
    """Tìm cột theo tên: khớp chính xác trước, sau đó khớp chuỗi con."""
    normed = [norm(c) for c in header_row]
    for cand in candidates:
        c = norm(cand)
        for i, h in enumerate(normed):
            if h and h == c:
                return i
    for cand in candidates:
        c = norm(cand)
        for i, h in enumerate(normed):
            if c and h and c in h:
                return i
    return None


def col_index(letter: str) -> int:
    """'A' -> 0"""
    return column_index_from_string(letter.strip().upper()) - 1


def get(row: list, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]
